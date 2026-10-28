"""Normalize official workbooks/API pages while preserving definitions and vintages."""
import argparse
import calendar
import csv
from collections import defaultdict
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
from decimal import Decimal

from openpyxl import load_workbook

from .geography import resolve, check_gst_code

MONTHS = {name.lower(): i for i, name in enumerate(calendar.month_name) if name}
MONTHS.update({name.lower(): i for i, name in enumerate(calendar.month_abbr) if name})


def monetary(value):
    if value is None or str(value).strip() in {'', '-', 'NA', 'N.A.', 'N/A', 'null'}:
        return None
    result = Decimal(str(value).replace(',', '').strip())
    if not result.is_finite() or result < 0:
        raise ValueError(f'Invalid value: {value}')
    return result


def month_header(value):
    if isinstance(value, datetime):
        return value.strftime('%Y-%m')
    match = re.fullmatch(r'([A-Za-z]{3})-(\d{2})', str(value))
    if match:
        return f'20{match[2]}-{MONTHS[match[1].lower()]:02d}'
    return None


def write_csv(path, rows):
    if not rows:
        raise ValueError(f'No records for {path}')
    fields = list(dict.fromkeys(k for row in rows for k in row))
    with Path(path).open('w', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, lineterminator='\n')
        writer.writeheader()
        writer.writerows(rows)


def base_row(state, month, source, meta, location, code=''):
    if source in {'gst', 'ewb'}:
        check_gst_code(code, state)
    geo_id, canonical, status = resolve(state, month, source)
    return {'state_id': geo_id, 'state': canonical, 'state_source': str(state).strip(),
            'state_code_source': code, 'geography_status': status,
            'reference_month': month, 'source': source, 'source_file': meta['file'],
            'source_url': meta['url'], 'source_sha256': meta['sha256'],
            'source_location': location, 'retrieved_at': meta['retrieved_at'],
            # Strict as-of joins use the earliest time this exact vintage is known.
            'available_at': meta['retrieved_at'][:10],
            'availability_basis': 'retrieval_only', 'synthetic': False}


def parse_gst(path, meta):
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook['Collections-Statewise']
    rows = list(sheet.values)
    if not any('State Wise Domestic GST Collection' in str(c) for row in rows[:6] for c in row):
        raise ValueError(f'Unverified GST definition in {path}')
    header = next(i for i, row in enumerate(rows) if str(row[1]).strip() == 'State')
    blocks = [(i, month_header(v)) for i, v in enumerate(rows[header]) if month_header(v)]
    if not blocks:
        raise ValueError('No monthly GST columns')
    output, checks = [], []
    for start, month in blocks:
        labels = []
        for label in rows[header + 1][start:]:
            labels.append(str(label).strip())
            if label == 'TOTAL':
                break
        if labels not in [['CGST', 'SGST', 'IGST', 'CESS', 'TOTAL'], ['CGST', 'SGST', 'IGST', 'TOTAL']]:
            raise ValueError(f'Unrecognized component layout: {labels}')
        summed = Decimal(0)
        published = None
        for idx, row in enumerate(rows[header + 2:], header + 3):
            if isinstance(row[0], (int, float)) or str(row[0]).isdigit():
                vals = {label: monetary(row[start + offset]) for offset, label in enumerate(labels)}
                obs = base_row(row[1], month, 'gst', meta, f'{sheet.title}!row{idx}', int(row[0]))
                components = [vals[x] for x in labels if x != 'TOTAL']
                if vals['TOTAL'] is not None:
                    summed += vals['TOTAL']
                if all(v is not None for v in components) and vals['TOTAL'] is not None:
                    if abs(sum(components) - vals['TOTAL']) > Decimal('0.05'):
                        raise ValueError(f'GST components do not reconcile: {path.name} {month} {row[1]}')
                obs.update({k.lower() + '_crore': '' if v is None else str(v) for k, v in vals.items()})
                obs['cess_in_published_total'] = 'included' if 'CESS' in labels else 'not_in_layout'
                obs['measure'] = 'domestic_gross_as_published'
                obs['gst_excluding_cess_crore'] = str(sum(vals[k] for k in ['CGST', 'SGST', 'IGST'])) if all(vals[k] is not None for k in ['CGST', 'SGST', 'IGST']) else ''
                output.append(obs)
            elif str(row[0]).strip() in {'Grand Total', 'Domestic GST Collection- All India'}:
                published = monetary(row[start + labels.index('TOTAL')])
        if published is None or abs(summed - published) > Decimal('0.1'):
            raise ValueError(f'GST state total reconciliation failed: {path.name} {month}')
        checks.append({'file': path.name, 'month': month, 'difference_crore': str(summed - published)})
    workbook.close()
    return output, checks


def parse_ewb(path, meta):
    workbook = load_workbook(path, data_only=True, read_only=True)
    output = []
    for sheet in workbook:
        rows = list(sheet.values)
        header = next(i for i, row in enumerate(rows) if len(row) >= 4 and str(row[1]).strip().upper() == 'STATE NAME')
        groups = [str(rows[header][i]).strip().upper() for i in [4, 7, 10]]
        if groups not in [['INTRA STATE SUPPLIES', 'INTER STATE OUTWARD SUPPLIES', 'INTER STATE INWARD SUPPLIES'], ['WITHIN-STATE', 'OUTGOING TO OTHER STATES', 'INCOMING FROM OTHER STATES']]:
            raise ValueError(f'Unknown EWB supply layout: {groups}')
        for start in [4, 7, 10]:
            labels = [str(v).upper() for v in rows[header + 1][start:start + 3]]
            if 'SUPPLIER' not in labels[0] or not any(token in labels[1] for token in ['EWAY', 'EWB']) or 'VALUE' not in labels[2] or 'CR' not in labels[2]:
                raise ValueError(f'Unknown EWB metric order/units: {labels}')
        for idx, row in enumerate(rows[header + 2:], header + 3):
            if not (isinstance(row[0], (int, float)) or str(row[0]).isdigit()):
                continue
            year, month_num = int(row[2]), int(row[3])
            month = datetime(year, month_num, 1).strftime('%Y-%m')
            obs = base_row(row[1], month, 'ewb', meta, f'{sheet.title}!row{idx}', int(row[0]))
            for start, supply in [(4, 'intra'), (7, 'outward'), (10, 'inward')]:
                for offset, metric in enumerate(['suppliers', 'bills', 'assessable_value_crore']):
                    value = monetary(row[start + offset])
                    if value is not None and metric != 'assessable_value_crore' and value != int(value):
                        raise ValueError('Fractional EWB count')
                    obs[f'ewb_{supply}_{metric}'] = '' if value is None else str(value)
            output.append(obs)
    workbook.close()
    return output


def parse_cpi(path, meta, state_codes):
    payload = json.loads(path.read_text())
    if payload.get('statusCode') is not True:
        raise ValueError(f'Invalid CPI response: {path}')
    output = []
    for idx, item in enumerate(payload['data']):
        base = str(item.get('baseyear', item.get('base_year')))
        if item['sector'] != 'Combined' or (item.get('group') != 'General' if base == '2012' else item.get('division') != 'CPI (General)'):
            raise ValueError('Non-headline or non-Combined CPI row')
        month = f'{int(item["year"]):04d}-{MONTHS[item["month"].lower()]:02d}'
        if item['state'] not in state_codes:
            raise ValueError(f'CPI state not in source metadata: {item["state"]}')
        obs = base_row(item['state'], month, f'cpi_{base}', meta, f'data[{idx}]', state_codes[item['state']])
        value = monetary(item.get('index'))
        obs.update({'cpi_index': '' if value is None else str(value), 'cpi_base': base,
                    'cpi_series': item.get('series', 'Current'),
                    'cpi_status': item.get('status', 'not_provided'),
                    'cpi_inflation_published': item.get('inflation')})
        output.append(obs)
    return output


def normalize(raw, output):
    raw, output = Path(raw), Path(output)
    output.mkdir(parents=True, exist_ok=True)
    if any(output.iterdir()):
        raise ValueError('Use an empty processed directory')
    manifest = json.loads((raw / 'manifest.json').read_text())
    validate_archive(raw, manifest)
    data = {'gst': [], 'ewb': [], 'cpi': []}
    checks = []
    cpi_codes = {}
    for base, series in [('2012', 'Current'), ('2024', 'Current'), ('2024', 'Back')]:
        filters = json.loads((raw / f'cpi_filters_{base}_{series}.json').read_text())['data']
        if isinstance(filters, list):
            filters = filters[0]
        cpi_codes[(base, series)] = {r['state_name']: r['state_code'] for r in filters['state']}
    for meta in manifest:
        path = raw / meta['file']
        if hashlib.sha256(path.read_bytes()).hexdigest() != meta['sha256']:
            raise ValueError(f'Source checksum changed: {path}')
        if path.name.startswith('statewise_'):
            parsed, validations = parse_gst(path, meta)
            data['gst'].extend(parsed); checks.extend(validations)
        elif path.name.startswith('ewb-data-'):
            data['ewb'].extend(parse_ewb(path, meta))
        elif re.fullmatch(r'cpi_\d{4}_(?:Current|Back)_\d+\.json', path.name):
            _, base, series, _ = path.stem.split('_')
            data['cpi'].extend(parse_cpi(path, meta, cpi_codes[(base, series)]))
    for source, rows in data.items():
        seen = set()
        for row in rows:
            identity = (row['state_id'], row['reference_month'], row.get('cpi_base', ''))
            if identity in seen:
                raise ValueError(f'Duplicate source observation: {source} {identity}')
            seen.add(identity)
        rows.sort(key=lambda r: (r['state_id'], r['reference_month'], r.get('cpi_base', '')))
        write_csv(output / f'{source}.csv', rows)
    (output / 'reconciliation.json').write_text(json.dumps(checks, indent=2))
    mappings = defaultdict(list)
    for rows in data.values():
        for row in rows:
            key = tuple(row[k] for k in ['source', 'state_code_source', 'state_source',
                                         'state_id', 'state', 'geography_status'])
            mappings[key].append(row['reference_month'])
    mapping_rows = []
    for key, months in sorted(mappings.items(), key=lambda pair: str(pair[0])):
        row = dict(zip(['source', 'source_code', 'source_name', 'state_id', 'state', 'status'], key))
        row.update({'first_month': min(months), 'last_month': max(months), 'observations': len(months)})
        mapping_rows.append(row)
    write_csv(output / 'state_mapping_audit.csv', mapping_rows)
    return data


def validate_archive(raw, manifest):
    """Reject partial downloads, including API success responses missing a page."""
    names = [entry['file'] for entry in manifest]
    if len(names) != len(set(names)):
        raise ValueError('Duplicate manifest filenames')
    html = (raw / 'gst_statistics_index.html').read_text()
    expected_files = set(re.findall(r'(?:statewise_GST_collection_|ewb-data-)[^"/]+\.xlsx', html))
    if not expected_files or not expected_files <= set(names):
        raise ValueError('Missing linked GST/EWB exports')
    for base, series in [('2012', 'Current'), ('2024', 'Current'), ('2024', 'Back')]:
        pages = [json.loads((raw / name).read_text()) for name in names
                 if re.fullmatch(f'cpi_{base}_{series}_[0-9]+\\.json', name)]
        if not pages or any(p.get('statusCode') is not True for p in pages):
            raise ValueError(f'Missing/failed CPI source: {base} {series}')
        metadata = pages[0]['meta_data']
        if {int(p['meta_data']['page']) for p in pages} != set(range(1, metadata['totalPages'] + 1)):
            raise ValueError('Missing CPI page')
        if any(p['meta_data']['totalRecords'] != metadata['totalRecords'] for p in pages):
            raise ValueError('CPI pagination changed')
        if sum(len(p['data']) for p in pages) != metadata['totalRecords']:
            raise ValueError('Incomplete CPI export')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--raw', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    result = normalize(args.raw, args.output)
    print({name: len(rows) for name, rows in result.items()})
